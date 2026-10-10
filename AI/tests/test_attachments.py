"""Attachments, document reading and voice input."""
from __future__ import annotations

import io
from unittest.mock import patch

import docx
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook
from PIL import Image

from AI.models import (
    AccountStatus,
    AttachmentKind,
    Conversation,
    Message,
    MessageAttachment,
    User,
)
from AI.services import (
    chat_service,
    model_service,
    settings_service,
    token_service,
    wallet_service,
)
from AI.tests.helpers import GPT, balance, fund, model_for
from AI.services.ai_service import AIError, AIResponse, ChatTurn, Transcription
from AI.utils.attachments import prepare, prepare_many
from AI.utils.extract import extract_text
from AI.utils.uploads import InvalidUpload

PASSWORD = "Sup3rSecret!x"


def png_bytes(colour="blue", size=(40, 40)) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, colour).save(buffer, format="PNG")
    return buffer.getvalue()


def docx_bytes(text="The mitochondria is the powerhouse of the cell.") -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Marks"
    sheet.append(["Student", "Score"])
    sheet.append(["Sokha", 88])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class ExtractionTests(TestCase):
    def test_word_documents_including_tables(self):
        document = docx.Document()
        document.add_paragraph("Intro line")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Stage"
        table.cell(0, 1).text = "Location"
        buffer = io.BytesIO()
        document.save(buffer)

        result = extract_text(buffer.getvalue(), "docx")

        self.assertIn("Intro line", result.text)
        self.assertIn("Stage | Location", result.text)

    def test_spreadsheets_keep_their_sheet_names(self):
        result = extract_text(xlsx_bytes(), "xlsx")
        self.assertIn("Sheet: Marks", result.text)
        self.assertIn("Sokha | 88", result.text)

    def test_csv_rows_become_readable_lines(self):
        result = extract_text(b"name,score\nSokha,88", "csv")
        self.assertEqual(result.text, "name | score\nSokha | 88")

    def test_a_corrupt_pdf_explains_itself_rather_than_raising(self):
        result = extract_text(b"definitely not a pdf", "pdf")
        self.assertFalse(result)
        self.assertIn("could not be opened", result.note)

    def test_long_files_are_truncated_with_a_note(self):
        result = extract_text(b"x" * 90_000, "txt")
        self.assertTrue(result.truncated)
        self.assertLessEqual(result.char_count, 24_000)
        self.assertTrue(result.note)

    def test_an_unknown_type_is_reported_not_crashed(self):
        self.assertIn("cannot be read", extract_text(b"data", "xyz").note)


class PrepareTests(TestCase):
    def test_an_image_is_recognised_and_normalised(self):
        item = prepare(SimpleUploadedFile("photo.png", png_bytes(), content_type="image/png"))
        self.assertEqual(item.kind, AttachmentKind.IMAGE)
        self.assertTrue(item.as_data_url().startswith("data:image/"))

    def test_a_document_is_read_at_upload_time(self):
        item = prepare(SimpleUploadedFile("notes.docx", docx_bytes()))
        self.assertEqual(item.kind, AttachmentKind.DOCUMENT)
        self.assertIn("powerhouse", item.text)

    def test_a_renamed_executable_is_refused(self):
        with self.assertRaises(InvalidUpload):
            prepare(SimpleUploadedFile("safe.pdf", b"MZ\x90\x00 windows binary"))

    def test_too_many_files_is_refused(self):
        files = [SimpleUploadedFile(f"n{i}.txt", b"hello there") for i in range(6)]
        with self.assertRaises(InvalidUpload) as caught:
            prepare_many(files, limit=5)
        self.assertIn("up to 5", str(caught.exception))

    @override_settings(UPLOAD_MAX_TOTAL_BYTES=1000)
    def test_the_combined_size_is_capped(self):
        files = [SimpleUploadedFile(f"n{i}.txt", b"x" * 600) for i in range(2)]
        with self.assertRaises(InvalidUpload) as caught:
            prepare_many(files, limit=5)
        self.assertIn("in total", str(caught.exception))

    def test_one_bad_file_fails_the_whole_batch(self):
        """A partly-accepted batch would bill for the wrong question."""
        files = [
            SimpleUploadedFile("good.txt", b"perfectly fine"),
            SimpleUploadedFile("bad.exe", b"MZ\x90\x00"),
        ]
        with self.assertRaises(InvalidUpload):
            prepare_many(files, limit=5)


class ChatAttachmentTests(TestCase):
    def setUp(self):
        settings_service.bootstrap()
        model_service.ensure_seeded()
        self.user = User.objects.create_user(
            username="attacher", email="attacher@test.local", password=PASSWORD,
            is_verified=True, status=AccountStatus.ACTIVE,
        )
        fund(self.user, 200_000, GPT)
        self.client.force_login(self.user)
        self.captured = {}

    def _fake_provider(self):
        captured = self.captured

        class Fake:
            name = "openai"
            model = "gpt-5-mini"
            is_configured = True
            supports_vision = True

            def complete(self, turns, system_prompt="", max_tokens=None, profile="standard"):
                captured["turns"] = list(turns)
                return AIResponse(
                    content="Answered.", input_tokens=500, output_tokens=100,
                    total_tokens=600, model="gpt-5-mini", request_id="r", latency_ms=50,
                )

        return Fake()

    def _send(self, **extra):
        payload = {"message": "Have a look at this"}
        payload.update(extra)
        with patch("AI.services.providers.for_model", return_value=self._fake_provider()):
            return self.client.post("/api/chat/send/", payload, HTTP_X_REQUESTED_WITH="XMLHttpRequest")

    def test_a_document_reaches_the_model_but_not_the_students_bubble(self):
        response = self._send(files=SimpleUploadedFile("biology.docx", docx_bytes()))
        self.assertEqual(response.status_code, 200)

        prompt = self.captured["turns"][-1].content
        self.assertIn("powerhouse", prompt)
        self.assertIn("BEGIN ATTACHED FILE: biology.docx", prompt)

        message = Message.objects.get(role="user")
        self.assertEqual(message.content, "Have a look at this")
        self.assertNotIn("powerhouse", message.content)

    def test_several_files_of_mixed_kinds_are_stored(self):
        response = self._send(files=[
            SimpleUploadedFile("a.docx", docx_bytes()),
            SimpleUploadedFile("b.png", png_bytes(), content_type="image/png"),
            SimpleUploadedFile("c.csv", b"x,y\n1,2"),
        ])
        self.assertEqual(response.status_code, 200)

        message = Message.objects.get(role="user")
        kinds = sorted(a.kind for a in message.attachments.all())
        self.assertEqual(kinds, ["document", "document", "image"])
        self.assertEqual(len(self.captured["turns"][-1].images()), 1)

    def test_every_image_is_passed_to_a_vision_model(self):
        self._send(files=[
            SimpleUploadedFile("one.png", png_bytes("red"), content_type="image/png"),
            SimpleUploadedFile("two.png", png_bytes("green"), content_type="image/png"),
        ])
        self.assertEqual(len(self.captured["turns"][-1].images()), 2)

    def test_a_follow_up_still_has_the_document_in_scope(self):
        self._send(files=SimpleUploadedFile("biology.docx", docx_bytes()))
        conversation = Conversation.objects.get(user=self.user)

        with patch("AI.services.providers.for_model", return_value=self._fake_provider()):
            self.client.post(
                "/api/chat/send/",
                {"message": "Which organelle was that?", "conversation_id": conversation.pk},
                HTTP_X_REQUESTED_WITH="XMLHttpRequest",
            )

        history = "\n".join(turn.content for turn in self.captured["turns"])
        self.assertIn("powerhouse", history)

    def test_document_text_is_paid_for(self):
        """A large attachment must be affordable before it is sent."""
        wallet_service.reclaim(
            self.user, model_for(GPT), balance(self.user, GPT) - 900, reason='Trim to 900'
        )
        response = self._send(files=SimpleUploadedFile("big.txt", b"word " * 4000))

        self.assertEqual(response.status_code, 402)
        self.assertEqual(response.json()["code"], "insufficient_tokens")

    def test_a_rejected_file_stops_the_whole_message(self):
        response = self._send(files=SimpleUploadedFile("virus.exe", b"MZ\x90\x00"))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["code"], "bad_upload")
        self.assertFalse(Message.objects.exists())

    def test_a_non_vision_model_is_told_in_words(self):
        class Blind:
            name = "openai"
            model = "text-only"
            is_configured = True
            supports_vision = False

            def complete(inner, turns, system_prompt="", max_tokens=None, profile="standard"):
                self.captured["turns"] = list(turns)
                return AIResponse(
                    content="ok", input_tokens=10, output_tokens=10, total_tokens=20,
                    model="text-only", request_id="r", latency_ms=10,
                )

        with patch("AI.services.providers.for_model", return_value=Blind()):
            self.client.post(
                "/api/chat/send/",
                {"message": "look", "files": SimpleUploadedFile("p.png", png_bytes(), content_type="image/png")},
                HTTP_X_REQUESTED_WITH="XMLHttpRequest",
            )

        self.assertIn("cannot read images", self.captured["turns"][-1].content)


class VoiceTests(TestCase):
    def setUp(self):
        settings_service.bootstrap()
        self.user = User.objects.create_user(
            username="speaker", email="speaker@test.local", password=PASSWORD,
            is_verified=True, status=AccountStatus.ACTIVE,
        )
        fund(self.user, 100_000, GPT)
        self.client.force_login(self.user)

    def _post(self, audio=b"AUDIO" * 400):
        return self.client.post(
            "/api/chat/transcribe/",
            {"audio": SimpleUploadedFile("speech.webm", audio, content_type="audio/webm")},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

    def test_the_text_comes_back_in_the_language_that_was_spoken(self):
        spoken = "मुझे प्रकाश संश्लेषण समझाइए"
        with patch(
            "AI.services.ai_service.transcribe",
            return_value=Transcription(text=spoken, language="hindi", duration=6.0, model="whisper-1"),
        ):
            data = self._post().json()

        self.assertTrue(data["ok"])
        self.assertEqual(data["text"], spoken)
        self.assertEqual(data["language"], "hindi")

    def test_usage_is_charged_by_duration(self):
        with patch(
            "AI.services.ai_service.transcribe",
            return_value=Transcription(text="hello", language="english", duration=60.0, model="whisper-1"),
        ):
            data = self._post().json()

        # One minute at the configured rate.
        self.assertEqual(data["tokens_charged"], 1500)

    def test_a_failed_transcription_costs_nothing(self):
        before = balance(self.user, GPT)
        with patch(
            "AI.services.ai_service.transcribe",
            side_effect=AIError("No speech was picked up.", code="no_speech", retryable=False),
        ):
            response = self._post()

        self.assertEqual(response.status_code, 502)
        self.assertEqual(balance(self.user, GPT), before)

    def test_an_empty_balance_blocks_voice_too(self):
        wallet_service.reclaim(
            self.user, model_for(GPT), balance(self.user, GPT), reason='Emptied'
        )
        response = self._post()
        self.assertEqual(response.status_code, 402)
        self.assertEqual(response.json()["code"], "no_tokens")

    def test_an_oversized_recording_is_refused_before_the_provider(self):
        with patch("AI.services.ai_service.transcribe") as called:
            response = self._post(audio=b"x" * (11 * 1024 * 1024))
        self.assertEqual(response.status_code, 400)
        called.assert_not_called()


class ConversationToolsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="tools", email="tools@test.local", password=PASSWORD,
            is_verified=True, status=AccountStatus.ACTIVE,
        )
        self.conversation = Conversation.objects.create(user=self.user, title="Revision plan")
        Message.objects.create(conversation=self.conversation, role="user", content="Help me revise")
        Message.objects.create(
            conversation=self.conversation, role="assistant", content="Start with **past papers**.",
            total_tokens=120, model_key="gpt-5-mini",
        )
        self.client.force_login(self.user)

    def test_archiving_hides_a_chat_without_deleting_it(self):
        response = self.client.post(
            f"/api/chat/{self.conversation.pk}/archive/", HTTP_X_REQUESTED_WITH="XMLHttpRequest"
        )
        self.assertTrue(response.json()["archived"])

        self.conversation.refresh_from_db()
        self.assertTrue(self.conversation.is_archived)
        self.assertEqual(self.conversation.messages.count(), 2, "archiving must not delete anything")

    def test_archiving_toggles_back(self):
        for _ in range(2):
            self.client.post(
                f"/api/chat/{self.conversation.pk}/archive/", HTTP_X_REQUESTED_WITH="XMLHttpRequest"
            )
        self.conversation.refresh_from_db()
        self.assertFalse(self.conversation.is_archived)

    def test_export_returns_a_markdown_download(self):
        response = self.client.get(f"/chat/{self.conversation.pk}/export/")

        self.assertEqual(response.status_code, 200)
        self.assertIn("text/markdown", response["Content-Type"])
        self.assertIn("attachment;", response["Content-Disposition"])

        body = response.content.decode()
        self.assertIn("# Revision plan", body)
        self.assertIn("Help me revise", body)
        self.assertIn("past papers", body)

    def test_another_student_cannot_export_it(self):
        stranger = User.objects.create_user(
            username="peeper", email="peeper@test.local", password=PASSWORD,
            is_verified=True, status=AccountStatus.ACTIVE,
        )
        self.client.force_login(stranger)
        self.assertEqual(self.client.get(f"/chat/{self.conversation.pk}/export/").status_code, 404)

    def test_another_student_cannot_archive_it(self):
        stranger = User.objects.create_user(
            username="meddler", email="meddler@test.local", password=PASSWORD,
            is_verified=True, status=AccountStatus.ACTIVE,
        )
        self.client.force_login(stranger)
        response = self.client.post(
            f"/api/chat/{self.conversation.pk}/archive/", HTTP_X_REQUESTED_WITH="XMLHttpRequest"
        )
        self.assertEqual(response.status_code, 404)
        self.conversation.refresh_from_db()
        self.assertFalse(self.conversation.is_archived)


class NotificationPanelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="reader", email="reader@test.local", password=PASSWORD,
            is_verified=True, status=AccountStatus.ACTIVE,
        )
        self.client.force_login(self.user)

    def test_the_feed_returns_this_users_notifications_only(self):
        from AI.models import Notification

        Notification.objects.create(user=self.user, title="Yours")
        other = User.objects.create_user(username="someone", email="s@test.local", password=PASSWORD)
        Notification.objects.create(user=other, title="Theirs")

        data = self.client.get("/api/notifications/", HTTP_X_REQUESTED_WITH="XMLHttpRequest").json()

        titles = [row["title"] for row in data["notifications"]]
        self.assertEqual(titles, ["Yours"])

    def test_marking_all_read_clears_the_badge(self):
        from AI.models import Notification

        Notification.objects.create(user=self.user, title="One")
        Notification.objects.create(user=self.user, title="Two")

        response = self.client.post("/api/notifications/read-all/", HTTP_X_REQUESTED_WITH="XMLHttpRequest")

        self.assertEqual(response.json()["unread"], 0)
        self.assertEqual(Notification.objects.filter(user=self.user, is_read=False).count(), 0)
